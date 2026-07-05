"""Full export (guide §6.1–6.2): JSON + XLSX + shareable expiring links."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime, timedelta

from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from shared.db.models import ExportJob, ForecastPointRow, ForecastRun


def _completed_run(client: TestClient, db_session: Session) -> str:
    cid = client.post(
        "/connectors",
        json={
            "connector_definition_key": "generic_rest",
            "name": "demo",
            "ingestion_method": "pull",
            "config": {"base_url": "http://x", "value_path": "$.v"},
        },
    ).json()["id"]
    mid = client.post(f"/connectors/{cid}/metrics", json={"name": "m", "key": "k"}).json()["id"]
    run_id = client.post(f"/metrics/{mid}/forecast", json={"horizon": 3}).json()["id"]
    run = db_session.get(ForecastRun, uuid.UUID(run_id))
    assert run is not None
    run.status = "completed"
    run.completed_at = datetime.now(tz=UTC)
    base = datetime(2026, 7, 1, tzinfo=UTC)
    for i in range(3):
        db_session.add(
            ForecastPointRow(
                forecast_run_id=run.id,
                timestamp=base + timedelta(minutes=i),
                predicted_value=10.0 + i,
                lower_bound=9.0 + i,
                upper_bound=11.0 + i,
            )
        )
    db_session.flush()
    return run_id


def test_json_export(client: TestClient, db_session: Session) -> None:
    run_id = _completed_run(client, db_session)
    resp = client.get(f"/forecasts/{run_id}/export?format=json")
    assert resp.status_code == 200
    body = resp.json()
    assert body["forecast_run_id"] == run_id
    assert len(body["points"]) == 3
    assert body["points"][0]["forecast"] == 10.0


def test_xlsx_export(client: TestClient, db_session: Session) -> None:
    from openpyxl import load_workbook

    run_id = _completed_run(client, db_session)
    resp = client.get(f"/forecasts/{run_id}/export?format=xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Forecast"]
    header = [c.value for c in ws[1]]
    assert header == ["timestamp", "actual", "forecast", "lower", "upper"]
    assert ws.max_row == 4  # header + 3 forecast rows
    assert "Model" in wb.sheetnames


def test_every_export_records_a_job(client: TestClient, db_session: Session) -> None:
    run_id = _completed_run(client, db_session)
    client.get(f"/forecasts/{run_id}/export?format=csv")
    jobs = db_session.scalars(
        select(ExportJob).where(ExportJob.forecast_run_id == uuid.UUID(run_id))
    ).all()
    assert len(jobs) == 1
    assert jobs[0].status == "completed"
    assert jobs[0].share_token
    assert jobs[0].expires_at is not None


def test_share_link_roundtrip(client: TestClient, db_session: Session) -> None:
    run_id = _completed_run(client, db_session)
    share = client.post(f"/forecasts/{run_id}/share?format=json").json()
    assert share["share_token"] in share["share_url"]

    # Public download by token — no auth, identical bytes.
    dl = client.get(f"/exports/{share['share_token']}")
    assert dl.status_code == 200
    assert dl.json()["forecast_run_id"] == run_id


def test_share_link_expires(client: TestClient, db_session: Session) -> None:
    run_id = _completed_run(client, db_session)
    token = client.post(f"/forecasts/{run_id}/share?format=csv").json()["share_token"]
    job = db_session.scalar(select(ExportJob).where(ExportJob.share_token == token))
    assert job is not None
    job.expires_at = datetime.now(tz=UTC) - timedelta(seconds=1)
    db_session.flush()
    assert client.get(f"/exports/{token}").status_code == 410


def test_unknown_share_token_404(client: TestClient) -> None:
    assert client.get("/exports/nope").status_code == 404
