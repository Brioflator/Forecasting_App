"""Export of a forecast in CSV / JSON / XLSX (guide §6.1–6.2), plus shareable
expiring links backed by export_jobs.share_token + BlobStore.signed_url.

Small forecasts generate synchronously and stream; every export also records an
export_jobs row and stores the artifact via the BlobStore so the share link
serves the identical bytes. The async large-export path is deliberately later
(guide §6.2 sizes it for hosted function limits, not local).
"""

from __future__ import annotations

import csv
import io
import json
import secrets
import uuid
from datetime import UTC, datetime, timedelta

from sqlalchemy import select
from sqlalchemy.orm import Session

from shared.db.models import DataPoint, ExportJob, ForecastPointRow, ForecastRun
from shared.providers.blob_store import BlobStore

SHARE_LINK_TTL = timedelta(days=7)

CONTENT_TYPES = {
    "csv": "text/csv",
    "json": "application/json",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _rows(session: Session, run: ForecastRun) -> list[dict]:
    """Chronological rows joining actuals and forecast on timestamp."""
    actuals = {
        r.timestamp: r.value
        for r in session.scalars(
            select(DataPoint).where(DataPoint.metric_id == run.metric_id)
        ).all()
    }
    forecast_rows = session.scalars(
        select(ForecastPointRow)
        .where(ForecastPointRow.forecast_run_id == run.id)
        .order_by(ForecastPointRow.timestamp)
    ).all()
    forecast_ts = {r.timestamp for r in forecast_rows}

    out: list[dict] = [
        {"timestamp": ts, "actual": actuals[ts], "forecast": None, "lower": None, "upper": None}
        for ts in sorted(t for t in actuals if t not in forecast_ts)
    ]
    for r in forecast_rows:
        out.append(
            {
                "timestamp": r.timestamp,
                "actual": actuals.get(r.timestamp),
                "forecast": r.predicted_value,
                "lower": r.lower_bound,
                "upper": r.upper_bound,
            }
        )
    return out


def render_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["timestamp", "actual", "forecast", "lower", "upper"])
    for row in rows:
        writer.writerow(
            [row["timestamp"].isoformat()]
            + [
                "" if row[k] is None else repr(row[k])
                for k in ("actual", "forecast", "lower", "upper")
            ]
        )
    return buf.getvalue().encode("utf-8")


def render_json(rows: list[dict], run: ForecastRun) -> bytes:
    body = {
        "forecast_run_id": str(run.id),
        "metric_id": str(run.metric_id),
        "model_params": run.model_params,
        "horizon": run.horizon,
        "points": [
            {
                "timestamp": row["timestamp"].isoformat(),
                "actual": row["actual"],
                "forecast": row["forecast"],
                "lower": row["lower"],
                "upper": row["upper"],
            }
            for row in rows
        ],
    }
    return json.dumps(body, indent=2).encode("utf-8")


def render_xlsx(rows: list[dict], run: ForecastRun) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.worksheet.worksheet import Worksheet

    wb = Workbook()
    ws = wb.active
    assert isinstance(ws, Worksheet)
    ws.title = "Forecast"
    header = ["timestamp", "actual", "forecast", "lower", "upper"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(
            [row["timestamp"].replace(tzinfo=None)]
            + [row[k] for k in ("actual", "forecast", "lower", "upper")]
        )
    ws.column_dimensions["A"].width = 22
    meta = wb.create_sheet("Model")
    meta.append(["forecast_run_id", str(run.id)])
    meta.append(["horizon", run.horizon])
    meta.append(["model_params", json.dumps(run.model_params)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


RENDERERS = {
    "csv": lambda rows, run: render_csv(rows),
    "json": render_json,
    "xlsx": render_xlsx,
}


def create_export(
    session: Session,
    blob_store: BlobStore,
    run: ForecastRun,
    fmt: str,
    requested_by: uuid.UUID,
) -> tuple[ExportJob, bytes]:
    """Render, store via the BlobStore, and record the export_jobs row.

    Returns (job, payload) — the endpoint streams payload directly and the
    share link (job.share_token) serves the stored copy later.
    """
    rows = _rows(session, run)
    payload = RENDERERS[fmt](rows, run)

    share_token = secrets.token_urlsafe(24)
    path = f"{run.organization_id}/{run.id}/{share_token}.{fmt}"
    blob_store.put(path, payload, CONTENT_TYPES[fmt])

    job = ExportJob(
        organization_id=run.organization_id,
        requested_by=requested_by,
        forecast_run_id=run.id,
        format=fmt,
        status="completed",
        file_path=path,
        share_token=share_token,
        expires_at=datetime.now(tz=UTC) + SHARE_LINK_TTL,
    )
    session.add(job)
    session.flush()
    return job, payload


# Back-compat shim for the POC surface (used by tests / forecasts router).
def forecast_to_csv(session: Session, run: ForecastRun) -> bytes:
    return render_csv(_rows(session, run))
